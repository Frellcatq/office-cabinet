# -*- coding: utf-8 -*-
"""Загрузка исходного XLSX (листы «Задачи», «Встречи», «Проекты», «Сотрудники»)
с полной проверкой перед заменой данных.

При любой ошибке функция возвращает список ошибок, и вызывающий код
оставляет текущие данные сеанса без изменений.
"""
from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .models import Employee, Meeting, Project, Task

SHEET_COLUMNS = {
    "Сотрудники": {"Код сотрудника", "ФИО", "Должность", "Подразделение"},
    "Проекты": {"Код проекта", "Название проекта", "Код владельца", "Начало",
                "Завершение", "Статус", "Цель проекта"},
    "Задачи": {"Код задачи", "Код проекта", "Название задачи", "Код ответственного",
               "Статус", "Приоритет", "Дата плана", "Начало работы", "План, ч",
               "Срок", "Дата завершения", "Комментарий"},
    "Встречи": {"Код встречи", "Код проекта", "Тема встречи", "Дата", "Начало",
                "Окончание", "Коды участников", "Место", "Статус", "Повестка"},
}


def _cell(row: Tuple, index: Dict[str, int], column: str):
    position = index.get(column)
    if position is None or position >= len(row):
        return None
    return row[position]


def _as_date(value, where: str, errors: List[str]) -> Optional[date]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    errors.append(f"{where}: не удалось прочитать дату «{value}».")
    return None


def _as_time(value, where: str, errors: List[str]) -> Optional[time]:
    if value in (None, ""):
        return None
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(text, fmt).time()
            except ValueError:
                pass
    errors.append(f"{where}: не удалось прочитать время «{value}».")
    return None


def _as_float(value, where: str, errors: List[str]) -> Optional[float]:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", "."))
        except ValueError:
            pass
    errors.append(f"{where}: не удалось прочитать число «{value}».")
    return None


def _code(value) -> str:
    return str(value).strip().upper() if value not in (None, "") else ""


def _text(value) -> str:
    return str(value).strip() if value not in (None, "") else ""


def parse_workbook(data: bytes) -> Tuple[Dict[str, list], List[str]]:
    """Разбирает книгу и проверяет столбцы, даты, коды и ссылки.

    Возвращает (наборы данных, список ошибок). Пустой набор данных
    при непустом списке ошибок означает «текущие данные не менять».
    """
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception:
        return {}, ["Не удалось открыть файл: ожидается книга XLSX."]

    errors: List[str] = []
    missing = [name for name in SHEET_COLUMNS if name not in workbook.sheetnames]
    if missing:
        return {}, ["В книге нет листов: " + ", ".join(missing) + "."]

    sheets: Dict[str, Tuple[list, Dict[str, int]]] = {}
    for name, required in SHEET_COLUMNS.items():
        rows = list(workbook[name].iter_rows(values_only=True))
        if not rows:
            errors.append(f"Лист «{name}» пуст.")
            continue
        header = [_text(h) for h in rows[0]]
        absent = required - set(header)
        if absent:
            errors.append(
                f"Лист «{name}»: нет столбцов {', '.join(sorted(absent))}."
            )
        sheets[name] = (rows[1:], {title: i for i, title in enumerate(header)})

    if errors:
        return {}, errors

    employees: List[Employee] = []
    projects: List[Project] = []
    tasks: List[Task] = []
    meetings: List[Meeting] = []

    # --- Сотрудники ---
    rows, index = sheets["Сотрудники"]
    for number, row in enumerate(rows, start=2):
        code = _code(_cell(row, index, "Код сотрудника"))
        if not code:
            continue
        employees.append(Employee(
            code=code,
            name=_text(_cell(row, index, "ФИО")),
            position=_text(_cell(row, index, "Должность")),
            department=_text(_cell(row, index, "Подразделение")),
        ))

    # --- Проекты ---
    rows, index = sheets["Проекты"]
    for number, row in enumerate(rows, start=2):
        code = _code(_cell(row, index, "Код проекта"))
        if not code:
            continue
        where = f"Проекты, строка {number} ({code})"
        projects.append(Project(
            code=code,
            name=_text(_cell(row, index, "Название проекта")),
            owner_code=_code(_cell(row, index, "Код владельца")),
            start=_as_date(_cell(row, index, "Начало"), where, errors),
            end=_as_date(_cell(row, index, "Завершение"), where, errors),
            status=_text(_cell(row, index, "Статус")) or "В работе",
            goal=_text(_cell(row, index, "Цель проекта")),
        ))

    # --- Задачи ---
    rows, index = sheets["Задачи"]
    for number, row in enumerate(rows, start=2):
        code = _code(_cell(row, index, "Код задачи"))
        if not code:
            continue
        where = f"Задачи, строка {number} ({code})"
        tasks.append(Task(
            code=code,
            project_code=_code(_cell(row, index, "Код проекта")),
            title=_text(_cell(row, index, "Название задачи")),
            assignee_code=_code(_cell(row, index, "Код ответственного")),
            status=_text(_cell(row, index, "Статус")) or "Новая",
            priority=_text(_cell(row, index, "Приоритет")) or "Средний",
            planned_date=_as_date(_cell(row, index, "Дата плана"), where, errors),
            start_time=_as_time(_cell(row, index, "Начало работы"), where, errors),
            planned_hours=_as_float(_cell(row, index, "План, ч"), where, errors),
            deadline=_as_date(_cell(row, index, "Срок"), where, errors),
            completed_date=_as_date(_cell(row, index, "Дата завершения"), where, errors),
            comment=_text(_cell(row, index, "Комментарий")),
        ))

    # --- Встречи ---
    rows, index = sheets["Встречи"]
    for number, row in enumerate(rows, start=2):
        code = _code(_cell(row, index, "Код встречи"))
        if not code:
            continue
        where = f"Встречи, строка {number} ({code})"
        participants_raw = _text(_cell(row, index, "Коды участников"))
        participants = [c for c in (p.strip().upper() for p in participants_raw.split(";")) if c]
        meetings.append(Meeting(
            code=code,
            project_code=_code(_cell(row, index, "Код проекта")),
            topic=_text(_cell(row, index, "Тема встречи")),
            date=_as_date(_cell(row, index, "Дата"), where, errors),
            start=_as_time(_cell(row, index, "Начало"), where, errors),
            end=_as_time(_cell(row, index, "Окончание"), where, errors),
            participant_codes=participants,
            place=_text(_cell(row, index, "Место")),
            status=_text(_cell(row, index, "Статус")) or "Запланирована",
            agenda=_text(_cell(row, index, "Повестка")),
        ))

    # --- Уникальность кодов ---
    def check_unique(codes: List[str], label: str) -> None:
        seen = set()
        for code in codes:
            if code in seen:
                errors.append(f"Код {code} встречается на листе «{label}» больше одного раза.")
            seen.add(code)

    check_unique([e.code for e in employees], "Сотрудники")
    check_unique([p.code for p in projects], "Проекты")
    check_unique([t.code for t in tasks], "Задачи")
    check_unique([m.code for m in meetings], "Встречи")

    # --- Ссылки между записями ---
    employee_codes = {e.code for e in employees}
    project_codes = {p.code for p in projects}
    for project in projects:
        if project.owner_code and project.owner_code not in employee_codes:
            errors.append(
                f"Проект {project.code}: владелец {project.owner_code} не найден среди сотрудников."
            )
    for task in tasks:
        if task.project_code not in project_codes:
            errors.append(f"Задача {task.code}: проект {task.project_code} не найден.")
        if task.assignee_code and task.assignee_code not in employee_codes:
            errors.append(
                f"Задача {task.code}: ответственный {task.assignee_code} не найден."
            )
    for meeting in meetings:
        if meeting.project_code not in project_codes:
            errors.append(f"Встреча {meeting.code}: проект {meeting.project_code} не найден.")
        for participant in meeting.participant_codes:
            if participant not in employee_codes:
                errors.append(
                    f"Встреча {meeting.code}: участник {participant} не найден."
                )

    if errors:
        return {}, errors

    return {
        "employees": employees,
        "projects": projects,
        "tasks": tasks,
        "meetings": meetings,
    }, []
